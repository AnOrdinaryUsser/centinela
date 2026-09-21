"""Loads the JCyL Red de Espacios Naturales Protegidos (REN) attribute
table into the protected_areas table.

Source: en_cyl_ren_limites.xlsx — 38 protected areas with name, protection
category, declared area (hectares) and an info URL.

IMPORTANT LIMITATION: this spreadsheet has NO boundary geometries, only
attributes. Rows are loaded with geom = NULL, which means the "is this
detection inside a protected area" check (ST_Contains in
backend/src/services/spatialService.js) will simply never match until real
polygon boundaries are loaded for these same 38 areas (e.g. from IDECYL or
patrimonionatural.org). Until then this table is informational only.

Expected local file: etl/data/en_cyl_ren_limites.xlsx
"""

from pathlib import Path

import openpyxl

from common.db import get_engine
from sqlalchemy import text

SOURCE_FILE = Path(__file__).resolve().parent.parent / "data" / "en_cyl_ren_limites.xlsx"
TARGET_TABLE = "protected_areas"


def main():
    wb = openpyxl.load_workbook(SOURCE_FILE, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter)
    col_index = {name: i for i, name in enumerate(header)}

    rows = []
    for row in rows_iter:
        rows.append(
            {
                "site_code": row[col_index["c_sitecode"]],
                "name": row[col_index["n_espacio"]],
                "protection_category": row[col_index["n_figura"]],
                "declared_area_hectares": row[col_index["m_sup_ha"]],
                "info_url": row[col_index["url_info"]],
                "source_dataset": "JCyL - Red de Espacios Naturales Protegidos (REN, sin geometria)",
            }
        )

    engine = get_engine()
    insert_stmt = text(
        f"""
        INSERT INTO {TARGET_TABLE}
            (site_code, name, protection_category, declared_area_hectares, info_url, source_dataset, geom)
        VALUES
            (:site_code, :name, :protection_category, :declared_area_hectares, :info_url, :source_dataset, NULL)
        """
    )

    with engine.begin() as conn:
        conn.execute(text(f"DELETE FROM {TARGET_TABLE} WHERE source_dataset LIKE 'JCyL - Red de Espacios Naturales Protegidos%'"))
        conn.execute(insert_stmt, rows)

    print(f"Loaded {len(rows)} rows into {TARGET_TABLE} (WITHOUT geometry - see module docstring)")


if __name__ == "__main__":
    main()
